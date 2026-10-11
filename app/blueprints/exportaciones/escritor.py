"""Escritura de las hojas a Excel, CSV o ZIP de CSV. No conoce la base de datos."""

import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
COLOR_MARCA = '86499D'   # púrpura de ZOE
_INVALIDOS = re.compile(r'[\[\]:*?/\\]')
_FORMULA = re.compile(r'^(=|\+|@|-\S)')


def _limpio(v):
    """Normaliza un valor y evita que un texto se interprete como fórmula."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'Sí' if v else 'No'
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str) and _FORMULA.match(v):
        return "'" + v
    return v


def _nombre_hoja(nombre, usados):
    base = _INVALIDOS.sub('', nombre)[:31] or 'Hoja'
    nombre, n = base, 2
    while nombre.lower() in usados:
        sufijo = f' ({n})'
        nombre = base[:31 - len(sufijo)] + sufijo
        n += 1
    usados.add(nombre.lower())
    return nombre


def a_xlsx(hojas, portada=None):
    """portada: dict opcional {'titulo', 'institucion', 'filtros', 'generado_por'}."""
    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    encabezado = PatternFill('solid', start_color=COLOR_MARCA)

    if portada:
        ws = wb.create_sheet(_nombre_hoja('Portada', usados))
        ws['A1'] = portada['titulo']
        ws['A1'].font = Font(name=FUENTE, size=16, bold=True, color=COLOR_MARCA)
        ws['A2'] = portada.get('institucion') or ''
        ws['A2'].font = Font(name=FUENTE, size=12)
        fila = 4
        ws.cell(fila, 1, 'Filtros aplicados').font = Font(name=FUENTE, bold=True)
        for linea in portada['filtros']:
            fila += 1
            ws.cell(fila, 1, linea).font = Font(name=FUENTE)
        fila += 2
        ws.cell(fila, 1, 'Generado').font = Font(name=FUENTE, bold=True)
        ws.cell(fila, 2, datetime.now().strftime('%Y-%m-%d %H:%M')).font = Font(name=FUENTE)
        fila += 1
        ws.cell(fila, 1, 'Generado por').font = Font(name=FUENTE, bold=True)
        ws.cell(fila, 2, portada.get('generado_por') or '').font = Font(name=FUENTE)
        fila += 2
        ws.cell(fila, 1, 'Contenido').font = Font(name=FUENTE, bold=True)
        ws.cell(fila, 2, 'Filas').font = Font(name=FUENTE, bold=True)
        indice_fila = fila
        ws.column_dimensions['A'].width = 48
        ws.column_dimensions['B'].width = 18
    else:
        ws = None

    creadas = []
    for h in hojas:
        hoja = wb.create_sheet(_nombre_hoja(h['nombre'], usados))
        creadas.append((hoja.title, len(h['rows'])))
        hoja.append(h['cols'])
        for c in hoja[1]:
            c.font = Font(name=FUENTE, bold=True, color='FFFFFF')
            c.fill = encabezado
            c.alignment = Alignment(vertical='center', wrap_text=True)
        anchos = [len(str(c)) for c in h['cols']]
        for fila_datos in h['rows']:
            limpia = [_limpio(v) for v in fila_datos]
            hoja.append(limpia)
            for i, v in enumerate(limpia[:len(anchos)]):
                anchos[i] = max(anchos[i], min(len(str(v)), 60))
        for row in hoja.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=FUENTE)
                if isinstance(c.value, datetime):
                    c.number_format = 'yyyy-mm-dd hh:mm'
                elif isinstance(c.value, date):
                    c.number_format = 'yyyy-mm-dd'
        for i, a in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = max(10, min(a + 2, 60))
        hoja.freeze_panes = 'A2'
        if h['cols']:
            hoja.auto_filter.ref = f"A1:{get_column_letter(len(h['cols']))}{max(1, len(h['rows']) + 1)}"

    if ws is not None:
        for i, (titulo, n) in enumerate(creadas, start=1):
            ws.cell(indice_fila + i, 1, titulo).font = Font(name=FUENTE)
            ws.cell(indice_fila + i, 2, n).font = Font(name=FUENTE)
    if not wb.sheetnames:
        wb.create_sheet('Sin datos')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _csv_valor(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date):
        return v.isoformat()
    return _limpio(v)


def _csv_bytes(hoja):
    out = io.StringIO()
    w = csv.writer(out, delimiter=';', lineterminator='\r\n')
    w.writerow(hoja['cols'])
    for fila in hoja['rows']:
        w.writerow([_csv_valor(v) for v in fila])
    return out.getvalue().encode('utf-8-sig')   # BOM: Excel abre bien tildes y ñ


def a_csv(hojas, nombre_base):
    """Una hoja -> (bytes, .csv). Varias -> ZIP con un CSV por hoja."""
    if len(hojas) == 1:
        return io.BytesIO(_csv_bytes(hojas[0])), f'{nombre_base}.csv', 'text/csv'
    buf = io.BytesIO()
    usados = set()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for h in hojas:
            z.writestr(_nombre_hoja(h['nombre'], usados) + '.csv', _csv_bytes(h))
    buf.seek(0)
    return buf, f'{nombre_base}_csv.zip', 'application/zip'
